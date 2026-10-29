"""Genera ../estacion7.pptx a partir de ../estacion7.html.

Cada slide se captura a 2560 × 1440 con Chrome y se inserta como imagen a
sangre completa, con su título como texto alternativo y nota del orador.
El PowerPoint se ve igual que el HTML, pero su texto no es editable: para
cambiar contenido, editar build-deck.py y volver a correr este script.

Uso: python3 build-pptx.py  (requiere Google Chrome y python-pptx)
"""

import html
import re
import subprocess
import tempfile
from pathlib import Path

from pptx import Presentation
from pptx.util import Emu

CHROME = "/Applications/Google Chrome.app/Contents/MacOS/Google Chrome"
ROOT = Path(__file__).resolve().parent.parent
SOURCE = ROOT / "estacion7.html"
TARGET = ROOT / "estacion7.pptx"


def main():
    page = SOURCE.read_text(encoding="utf-8")
    page = page.replace('src="assets/', f'src="file://{ROOT}/assets/')
    titles = [html.unescape(t) for t in re.findall(r'aria-label="\d+\. ([^"]+)"', page)]
    modules = [m.strip().title() for m in re.findall(r"Estación 7 · ([^<]+)</span>", page)]

    pres = Presentation()
    pres.slide_width, pres.slide_height = Emu(12192000), Emu(6858000)
    blank = pres.slide_layouts[6]

    with tempfile.TemporaryDirectory() as tmp:
        tmp = Path(tmp)
        for n, (title, module) in enumerate(zip(titles, modules), 1):
            css = (
                "<style>.deck{padding:0!important;gap:0!important;display:block!important}"
                f".slide{{display:none}}.slide:nth-of-type({n}){{display:block}}"
                "body{margin:0;background:#0A0A0A}</style>"
            )
            single = tmp / f"s{n:02d}.html"
            single.write_text(page.replace("</head>", css + "</head>"), encoding="utf-8")
            png = tmp / f"s{n:02d}.png"
            subprocess.run(
                [CHROME, "--headless=new", "--disable-gpu", "--hide-scrollbars",
                 "--force-device-scale-factor=2", "--window-size=1280,720",
                 "--virtual-time-budget=3000", "--allow-file-access-from-files",
                 f"--screenshot={png}", single.as_uri()],
                check=True, capture_output=True,
            )
            slide = pres.slides.add_slide(blank)
            pic = slide.shapes.add_picture(str(png), 0, 0, pres.slide_width, pres.slide_height)
            pic.name = f"Slide {n}"
            pic._element.nvPicPr.cNvPr.set("descr", f"{title} ({module})")
            slide.notes_slide.notes_text_frame.text = (
                f"{module} · {title}\nGuía del laboratorio: estacion7/laboratorio.md"
            )

    pres.save(TARGET)
    print(f"{len(titles)} slides → {TARGET}")


if __name__ == "__main__":
    main()
